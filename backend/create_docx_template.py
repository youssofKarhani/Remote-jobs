"""Generate canonical Jinja2-compatible Word document template for CV generation."""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_template(output_path: Path):
    doc = docx.Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    # Candidate Name (Header)
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("{{ candidate_name }}")
    r_name.font.size = Pt(22)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800

    # Headline
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_head = p_head.add_run("{{ headline }}")
    r_head.font.size = Pt(12)
    r_head.font.italic = True
    r_head.font.color.rgb = RGBColor(0x47, 0x55, 0x69)  # Slate 600

    # Contact line
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_contact = p_contact.add_run("{{ email }}  |  {{ phone }}  |  {{ location }}")
    r_contact.font.size = Pt(9.5)
    r_contact.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Links line
    p_links = doc.add_paragraph()
    p_links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_links = p_links.add_run("{{ linkedin_url }}  |  {{ github_url }}  |  {{ portfolio_url }}")
    r_links.font.size = Pt(9)
    r_links.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6)  # Blue link color

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title.upper())
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)  # Slate 900
        # Add bottom border via separator
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(0)
        p_div.paragraph_format.space_after = Pt(4)
        r_div = p_div.add_run("_________________________________________________________________________________")
        r_div.font.size = Pt(6)
        r_div.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    # Professional Summary
    add_section_header("Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    r_sum = p_sum.add_run("{{ summary }}")
    r_sum.font.size = Pt(10)
    r_sum.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Technical Skills
    add_section_header("Skills & Technologies")
    p_skills = doc.add_paragraph()
    p_skills.paragraph_format.space_after = Pt(6)
    r_skills = p_skills.add_run("{% for skill in skills %}{{ skill.name }}{% if not loop.last %}  •  {% endif %}{% endfor %}")
    r_skills.font.size = Pt(9.5)
    r_skills.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Work Experience
    add_section_header("Work Experience")
    p_exp_loop_start = doc.add_paragraph("{% for exp in experiences %}")
    p_exp_loop_start.paragraph_format.space_after = Pt(2)
    
    p_exp_head = doc.add_paragraph()
    p_exp_head.paragraph_format.space_before = Pt(4)
    p_exp_head.paragraph_format.space_after = Pt(2)
    r_exp_title = p_exp_head.add_run("{{ exp.role_title }} — {{ exp.company_name }}")
    r_exp_title.font.bold = True
    r_exp_title.font.size = Pt(10.5)
    r_exp_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    p_exp_sub = doc.add_paragraph()
    p_exp_sub.paragraph_format.space_after = Pt(4)
    r_exp_sub = p_exp_sub.add_run("{{ exp.start_date }} - {{ exp.end_date }}  |  {{ exp.location }}")
    r_exp_sub.font.italic = True
    r_exp_sub.font.size = Pt(9)
    r_exp_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Bullet loop
    p_bullet_start = doc.add_paragraph("{% for bullet in exp.bullets %}")
    p_bullet_start.paragraph_format.space_after = Pt(1)
    
    p_bullet = doc.add_paragraph(style='List Bullet')
    p_bullet.paragraph_format.space_after = Pt(2)
    r_bullet = p_bullet.add_run("{{ bullet }}")
    r_bullet.font.size = Pt(9.5)
    r_bullet.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    p_bullet_end = doc.add_paragraph("{% endfor %}")
    p_bullet_end.paragraph_format.space_after = Pt(4)

    p_exp_loop_end = doc.add_paragraph("{% endfor %}")
    p_exp_loop_end.paragraph_format.space_after = Pt(6)

    # Projects
    add_section_header("Key Projects")
    doc.add_paragraph("{% for proj in projects %}")
    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_after = Pt(2)
    r_proj_t = p_proj.add_run("{{ proj.title }}  ")
    r_proj_t.font.bold = True
    r_proj_t.font.size = Pt(10)
    r_proj_tech = p_proj.add_run("({{ proj.technologies }})")
    r_proj_tech.font.italic = True
    r_proj_tech.font.size = Pt(9)
    
    p_proj_desc = doc.add_paragraph(style='List Bullet')
    p_proj_desc.paragraph_format.space_after = Pt(4)
    r_proj_desc = p_proj_desc.add_run("{{ proj.description }}")
    r_proj_desc.font.size = Pt(9.5)
    doc.add_paragraph("{% endfor %}")

    # Education
    add_section_header("Education")
    doc.add_paragraph("{% for edu in education %}")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_after = Pt(2)
    r_edu = p_edu.add_run("{{ edu.degree }} in {{ edu.field_of_study }} — {{ edu.institution }} ({{ edu.start_date }} - {{ edu.end_date }})")
    r_edu.font.size = Pt(9.5)
    doc.add_paragraph("{% endfor %}")

    # Certifications
    add_section_header("Certifications")
    doc.add_paragraph("{% for cert in certifications %}")
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_after = Pt(2)
    r_cert = p_cert.add_run("{{ cert.name }} — {{ cert.issuing_organization }} ({{ cert.issue_date }})")
    r_cert.font.size = Pt(9.5)
    doc.add_paragraph("{% endfor %}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Generated template at: {output_path}")

if __name__ == "__main__":
    out = Path("app/templates/cv/modern_template.docx")
    build_template(out)
